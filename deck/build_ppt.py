#!/usr/bin/env python3
"""Assemble the SIH26249 deck: 10 full-bleed slides from ./assets PNGs,
speaker notes on every slide, document properties set."""
import os
from pptx import Presentation
from pptx.util import Cm

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "SAJJATA_SIH26249_Idea_Submission.pptx")

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(26.67), Cm(15.0)
blank = prs.slide_layouts[6]

NOTES = [
# 1
"""SAJJATA (sajjata = readiness). One line: turn fragmented maintenance data into proactive decisions, and decisions into aircraft availability.
Open with the PS verbatim - availability is lost to fragmentation and reactive practice. We will show the full chain: data to prediction to approved work order to fleet readiness.
The problem statement publishes NO dataset; every number shown later that is not a target is measured by us on clearly-labelled synthetic data. Say this up front - credibility is a weapon in this room.""",
# 2
"""Four real data silos from the PS: health monitoring, technical records, spares, maintenance agencies. The right panel is the causal chain the PS itself states: delayed fault prediction, avoidable downtime, sub-optimal utilisation.
Positioning line: this is a data-integration failure first and an AI problem second. No fleet statistics are claimed - the panel restates the official problem chain only.""",
# 3
"""Seven verbs - Sense, Integrate, Predict, Explain, Optimise, Act, Learn. Emphasise three pillars: (1) multi-model AI, not one black box; (2) the digital twin as the live, auditable aircraft state; (3) every prediction ends in a human-approved work order - we do not stop at a probability.""",
# 4
"""Walk left to right: sources -> validated ingestion (streaming with store-and-forward, batch loaders, quality gates) -> data platform (time-series + PostgreSQL + object lake + feature store) -> intelligence (twin, model services, risk engine) -> decision & action (maintenance decision engine, fleet optimiser, spares, work orders).
Red band: security envelope wraps all layers - zero-trust, RBAC+MFA, TLS + AES-256, immutable audit, classification, segmentation, secrets, secure model registry.""",
# 5
"""Model A anomaly: isolation forest on per-unit z-scored sensors (each aircraft baselined against its own healthy period). Model B: boosted trees, unit-disjoint split. Model C: RUL regression. Model D routes severity.
Star means MEASURED on our synthetic benchmark (repo: deck/ml_benchmark.py) - methodology demonstration, not fleet performance. Right: A-017 twin showing the fuel-nozzle story; twin updates on telemetry, maintenance, replacement, inspection, prediction changes - versioned.""",
# 6
"""The 16-step closed loop. Pause on step 12: HUMAN APPROVAL - AI recommends, human decides, system learns. Outcomes loop back as training data. This is the slide that answers 'so what' - nothing in the chain dead-ends.""",
# 7
"""Fleet optimiser: inputs left, objective centre (maximise available aircraft subject to safety + maintenance constraints; greedy + constraint checks in prototype, OR-Tools solvers in production), this-week plan right.
Spares: prediction -> parts -> reservation -> procurement -> schedule alignment. Demo table shows a real shortage case (INS cartridge, 21-day lead) - the system flags it BEFORE the failure window.""",
# 8
"""Six screens, each ending in an action: commander's fleet picture, one-aircraft twin, the alert decision card (WHAT/WHEN/WHY/CONF/ACTION), planner with bays and crews, forecast-aware inventory, and analytics that grade the AI itself (precision/recall + drift monitor).""",
# 9
"""Three honest bands: IMPLEMENTED (end-to-end demo pipeline on synthetic data, reproducible benchmark), PROTOTYPE-READY (scale-out ingestion, connectors, solvers, SSO/MFA, HA), PRODUCTION/FUTURE (accredited MoD hosting, official system-of-record integration, safety case).
MLOps loop top-right: registry, versioned features, shadow evaluation, rollback, drift monitors. Phased path: P0 synthetic pilot (weeks 1-6), P1 shadow mode on one unit's data (weeks 7-16), P2 production integration + accreditation.""",
# 10
"""Targets, not claims: availability +3-5pp, unscheduled removals -20-30%, lead time >=7 days, MTTR -15-25%, critical stockouts ->0, alert precision >=80%. Each has a published measurement definition; nothing is presented as achieved.
Close on the line: FROM REACTIVE MAINTENANCE TO PREDICTIVE FLEET READINESS. References are public policy and benchmark sources (DoDI 4151.22, CBM+ Guidebook, USAF RSO-CBM+, NASA C-MAPSS, NIST SP 800-207, SHAP).""",
]

for i in range(10):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(os.path.join(ASSETS, f"s{i+1:02d}.png"), 0, 0,
                         width=prs.slide_width, height=prs.slide_height)
    s.notes_slide.notes_text_frame.text = NOTES[i]

cp = prs.core_properties
cp.title = "SAJJATA - Air Power: Predictive Maintenance & Fleet Availability (SIH26249)"
cp.subject = "Smart India Hackathon 2026 - PS SIH26249 - MoD / Defence Services Staff College"
cp.author = "Team <TEAM NAME>"
cp.keywords = "SIH26249, predictive maintenance, digital twin, fleet availability, PHM"
cp.comments = ("Synthetic-data prototype. Measured metrics are from deck/ml_benchmark.py on synthetic "
               "data; all operational numbers are labelled targets.")

prs.save(OUT)
print("saved", OUT)
